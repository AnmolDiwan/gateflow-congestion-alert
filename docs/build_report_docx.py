"""Builds docs/GateFlow_Report.docx from docs/REPORT.md (one source of truth for the report).
Handles the Markdown subset the report uses: headings, paragraphs, **bold** / *italic* / `code`,
bullet and numbered lists (one nesting level), tables, images, block quotes, code blocks, rules.
Usage: python docs/build_report_docx.py
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = Path(__file__).parent
SRC, OUT = HERE / "REPORT.md", HERE / "GateFlow_Report.docx"
INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")


def add_runs(par, text):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**"):
            par.add_run(part[2:-2]).bold = True
        elif part.startswith("`"):
            r = par.add_run(part[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(9.5)
        elif part.startswith("["):
            label, url = re.match(r"\[([^\]]+)\]\(([^)]+)\)", part).groups()
            par.add_run(f"{label} ({url})")
        elif part.startswith("*"):
            par.add_run(part[1:-1]).italic = True
        else:
            par.add_run(part)


def shade(cell, hex_fill):
    tc = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), hex_fill)
    tc.append(sh)


def add_table(doc, rows):
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    header, body = cells[0], [r for r in cells[2:]]
    t = doc.add_table(rows=1 + len(body), cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(header):
        c = t.cell(0, j)
        c.text = ""
        add_runs(c.paragraphs[0], f"**{h}**" if h else "")
        shade(c, "E8F1FB")
    for i, r in enumerate(body, start=1):
        for j in range(len(header)):
            c = t.cell(i, j)
            c.text = ""
            add_runs(c.paragraphs[0], r[j] if j < len(r) else "")
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9.5)
    doc.add_paragraph()


def main():
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    for s in ("Heading 1", "Heading 2", "Heading 3"):
        doc.styles[s].font.color.rgb = RGBColor(0x1C, 0x5C, 0xAB)
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Inches(0.9)

    lines = SRC.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        s = line.strip()
        if s.startswith("```"):
            i += 1
            code = []
            while not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            p = doc.add_paragraph()
            r = p.add_run("\n".join(code))
            r.font.name = "Consolas"
            r.font.size = Pt(9.5)
        elif s.startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            add_table(doc, block)
            continue
        elif s.startswith("!["):
            alt, path = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", s).groups()
            img = (HERE / path).resolve()
            if img.exists():
                doc.add_picture(str(img), width=Inches(6.5))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap = doc.add_paragraph(alt)
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap.runs[0].italic = True
        elif s.startswith("# "):
            doc.add_heading(s[2:], level=0)
        elif s.startswith("## "):
            doc.add_heading(s[3:], level=1)
        elif s.startswith("### "):
            doc.add_heading(s[4:], level=2)
        elif s.startswith("> "):
            p = doc.add_paragraph(style="Intense Quote")
            add_runs(p, s[2:])
        elif s == "---":
            pass
        elif re.match(r"^\s*- ", line):
            lvl = 2 if line.startswith("    ") else 1
            p = doc.add_paragraph(style="List Bullet" if lvl == 1 else "List Bullet 2")
            add_runs(p, re.sub(r"^\s*- ", "", line))
        elif re.match(r"^\s*\d+\. ", line):
            p = doc.add_paragraph(style="List Number")
            add_runs(p, re.sub(r"^\s*\d+\. ", "", line))
        elif s:
            p = doc.add_paragraph()
            add_runs(p, s)
        i += 1
    doc.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
